"""Text extraction from uploaded case materials (PDF, Word, plain text, ZIP bundles).

Class negotiation cases arrive as a handful of files: general instructions plus one
confidential sheet per role, sometimes with score tables. We only need their text; the
parser (parser.parse_documents) works out which file is which.
"""

from __future__ import annotations

import io
import stat
from collections.abc import Iterable, Iterator
from pathlib import Path, PurePosixPath
from zipfile import ZipFile

SUPPORTED = ("pdf", "docx", "txt", "md")
UPLOAD_TYPES = (*SUPPORTED, "zip")
MAX_ZIP_ENTRIES = 200
MAX_ZIP_BYTES = 50 * 1024 * 1024


class UnreadableDocument(ValueError):
    pass


def extract_documents(uploads: Iterable[tuple[str, bytes]]) -> tuple[dict[str, str], list[str]]:
    """Read loose files and ZIP bundles, keeping readable documents and reporting skips.

    Archive paths identify the source of each sheet. Repeated names get a suffix so
    uploading two files with the same name never silently drops one of them.
    """
    documents: dict[str, str] = {}
    warnings: list[str] = []
    for name, data in uploads:
        files = (
            _zip_files(name, data, warnings)
            if Path(name).suffix.lower() == ".zip"
            else [(name, data)]
        )
        for filename, content in files:
            try:
                text = extract_text(filename, content)
            except UnreadableDocument as e:
                warnings.append(str(e))
                continue
            label = filename
            number = 2
            while label in documents:
                label = f"{filename} ({number})"
                number += 1
            documents[label] = text
    return documents, warnings


def _zip_files(name: str, data: bytes, warnings: list[str]) -> Iterator[tuple[str, bytes]]:
    """Read members in memory; never extract archive paths onto the filesystem."""
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                warnings.append(f"{name}: too many ZIP entries (limit {MAX_ZIP_ENTRIES}).")
                return
            members = []
            for entry in entries:
                path = PurePosixPath(entry.filename.replace("\\", "/"))
                label = f"{name}/{path}"
                if entry.is_dir():
                    continue
                if path.is_absolute() or ".." in path.parts:
                    warnings.append(f"{label}: skipped unsafe archive path.")
                    continue
                if any(part.startswith(".") or part == "__MACOSX" for part in path.parts):
                    continue
                if stat.S_ISLNK(entry.external_attr >> 16):
                    warnings.append(f"{label}: skipped symbolic link.")
                    continue
                if path.suffix.lower().lstrip(".") not in SUPPORTED:
                    warnings.append(f"{label}: skipped unsupported file type.")
                    continue
                members.append((entry, label))
            if not members:
                warnings.append(f"{name}: no supported documents found in ZIP.")
                return
            if sum(entry.file_size for entry, _ in members) > MAX_ZIP_BYTES:
                warnings.append(
                    f"{name}: documents exceed the {MAX_ZIP_BYTES // (1024 * 1024)} MB "
                    "uncompressed ZIP limit. Upload a smaller bundle."
                )
                return
            remaining = MAX_ZIP_BYTES
            for entry, label in members:
                if entry.flag_bits & 1:
                    warnings.append(f"{label}: password-protected files are not supported.")
                    continue
                try:
                    with archive.open(entry) as member:
                        content = member.read(remaining + 1)
                except Exception as e:  # noqa: BLE001 - corrupt/unsupported ZIP members
                    warnings.append(f"{label}: could not read ZIP entry: {e}")
                    continue
                remaining -= len(content)
                if remaining < 0:
                    warnings.append(f"{name}: uncompressed ZIP size limit exceeded.")
                    return
                yield label, content
    except Exception as e:  # noqa: BLE001 - malformed ZIP archives should not crash uploads
        warnings.append(f"{name}: could not read ZIP: {e}")


def extract_text(name: str, data: bytes) -> str:
    """Plain text of one uploaded file. Raises UnreadableDocument if nothing is extractable
    (for example a scanned PDF with no text layer)."""
    suffix = Path(name).suffix.lower().lstrip(".")
    if suffix in ("pdf", "docx"):
        try:
            text = _pdf(data) if suffix == "pdf" else _docx(data)
        except UnreadableDocument as e:
            raise UnreadableDocument(f"{name}: {e}") from e
    elif suffix in ("txt", "md"):
        text = data.decode("utf-8", errors="replace")
    else:
        raise UnreadableDocument(f"{name}: unsupported file type (use {', '.join(SUPPORTED)})")
    text = text.strip()
    if not text:
        raise UnreadableDocument(
            f"{name}: no text found (a scanned PDF? export it with a text layer or paste it)"
        )
    return text


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:  # noqa: BLE001 - pypdf raises many types on malformed files
        raise UnreadableDocument(f"could not read PDF: {e}") from e


def _docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise UnreadableDocument(f"could not read Word file: {e}") from e
    parts = [p.text for p in document.paragraphs]
    # Score tables in class cases are usually Word tables; keep them as tab-separated rows.
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)
