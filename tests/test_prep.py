import io

import docx
import pytest

from negotrain.cases.documents import UnreadableDocument, extract_text
from negotrain.cases.parser import parse_documents
from negotrain.llm import ScriptedLLM
from negotrain.prep import build_prep, prep_numbers, priorities
from test_parser import parsed_from


def make_pdf(text: str) -> bytes:
    """A minimal one-page PDF with a text layer."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_docx(paragraph: str, table: list[list[str]]) -> bytes:
    document = docx.Document()
    document.add_paragraph(paragraph)
    t = document.add_table(rows=len(table), cols=len(table[0]))
    for r, row in enumerate(table):
        for c, cell in enumerate(row):
            t.cell(r, c).text = cell
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_extract_text_from_pdf_docx_and_txt():
    assert "Confidential instructions" in extract_text(
        "buyer.pdf", make_pdf("Confidential instructions")
    )
    text = extract_text(
        "scores.docx", make_docx("Score sheet", [["Price", "Points"], ["$10", "40"]])
    )
    assert "Score sheet" in text and "Price\tPoints" in text
    assert extract_text("notes.md", b"# Public info") == "# Public info"


def test_extract_text_rejects_empty_and_unsupported_files():
    with pytest.raises(UnreadableDocument, match="no text"):
        extract_text("blank.txt", b"   ")
    with pytest.raises(UnreadableDocument, match="unsupported"):
        extract_text("deck.pptx", b"x")
    with pytest.raises(UnreadableDocument):
        extract_text("broken.pdf", b"not a pdf")


def test_parse_documents_sends_every_file_and_asks_to_infer(case):
    llm = ScriptedLLM([parsed_from(case)])
    got = parse_documents({"general.pdf": "public facts", "candidate.docx": "my sheet"}, llm)
    assert got.id == case.id
    content = llm.calls[0][1][0]["content"]
    assert "# File: general.pdf" in content and "my sheet" in content
    assert "infer a plausible sheet" in content


def test_priorities_rank_issues_by_points_at_stake(case):
    ranked = priorities(case, "candidate")
    assert [p.issue_id for p in ranked][:2] == ["salary", "remote"]
    assert ranked[0].best_option == "$150,000"
    assert abs(sum(p.share for p in ranked) - 1) < 1e-9


def test_prep_numbers(case):
    n = prep_numbers(case, "candidate")
    assert n.walk_away == 3500 and n.best_possible == 9900
    assert n.walk_away < n.target < n.best_possible


def test_template_prep_without_llm(case):
    prep = build_prep(case, "candidate")
    assert prep.brief.source == "template"
    assert prep.brief.checklist and "Base salary" in prep.brief.tradeoffs[0]


def test_llm_prep_never_sees_the_other_sides_sheet(case):
    canned = {
        "overview": "o",
        "key_facts": ["k"],
        "your_interests": ["y"],
        "their_likely_interests": ["they may..."],
        "tradeoffs": ["t"],
        "questions_to_ask": ["q?"],
        "opening_plan": "open",
        "pitfalls": ["p"],
        "checklist": ["c"],
    }
    llm = ScriptedLLM([canned])
    prep = build_prep(case, "candidate", llm)
    assert prep.brief.source == "llm"
    system, messages = llm.calls[0]
    prompt = system + messages[0]["content"]
    assert case.role("candidate").confidential_brief in prompt
    assert case.role("recruiter").confidential_brief not in prompt

    def boom(system, messages):
        raise RuntimeError("down")

    assert build_prep(case, "candidate", ScriptedLLM([boom])).brief.source == "template"
